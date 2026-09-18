"""Build the Chinese SAP with the bundled document runtime; no patient data."""
from __future__ import annotations

import csv
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
from wmh_hcy.studies.registry import SOURCES, STUDIES, UNITS, roles


def main():
    text = (ROOT / "docs/four_studies_plan.md").read_text(encoding="utf-8")
    doc = Document()
    # Bundled default templates can carry a decorative Title border. Remove it.
    for element in [doc.styles.element, doc.element]:
        for border in list(element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.75)
    section.left_margin = section.right_margin = Inches(.8)
    for name, size in [("Normal", 11), ("Title", 22), ("Heading 1", 16), ("Heading 2", 12)]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.16
    doc.styles["Normal"].paragraph_format.widow_control = True
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(0)
    header = section.header.paragraphs[0]
    header.text = "CNSR III  影像与长期预后  |  统计分析方案  |  2026年9月18日"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor(0, 0, 0)
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run("第 ").font.size = Pt(9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    footer.add_run(" 页").font.size = Pt(9)
    lines = text.splitlines()
    i = 0
    page_before = False
    while i < len(lines):
        line = lines[i].strip()
        i += 1
        if not line:
            continue
        if line == "<!-- PAGE -->":
            page_before = True
        elif line.startswith("# "):
            doc.add_paragraph(line[2:].replace("四项独立研究", "\n四项独立研究"), "Title")
        elif line.startswith("## "):
            paragraph = doc.add_paragraph(line[3:], "Heading 1")
            paragraph.paragraph_format.page_break_before = page_before
            page_before = False
        elif line.startswith("### "):
            doc.add_paragraph(line[4:], "Heading 2")
        elif line.startswith("|"):
            rows = [line]
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not re.match(r"^\|[\s|:-]+$", lines[i]):
                    rows.append(lines[i].strip())
                i += 1
            values = [[c.strip() for c in row.strip("|").split("|")] for row in rows]
            table = doc.add_table(rows=1, cols=len(values[0]))
            table.autofit = False
            table.style = "Table Grid"
            if len(values[0]) == 3:
                widths = [1.32, 2.7, 2.88]
            elif values[0][0] == "原始记录与规则证据":
                widths = [3.5, 3.4]
            else:
                widths = [1.45, 5.45]
            for j, width in enumerate(widths):
                table.columns[j].width = Inches(width)
            for ri, values_row in enumerate(values):
                cells = table.rows[0].cells if ri == 0 else table.add_row().cells
                for ci, value in enumerate(values_row):
                    cells[ci].width = Inches(widths[ci])
                    cells[ci].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    cells[ci].text = value
                    for p in cells[ci].paragraphs:
                        p.paragraph_format.space_after = Pt(5)
                        p.paragraph_format.space_before = Pt(5)
                        p.paragraph_format.line_spacing = 1.08
                        for r in p.runs:
                            r.font.size = Pt(10)
                            r.bold = ri == 0
                    if ri == 0:
                        shade = OxmlElement("w:shd")
                        shade.set(qn("w:fill"), "EFEFEF")
                        cells[ci]._tc.get_or_add_tcPr().append(shade)
                trpr = table.rows[ri]._tr.get_or_add_trPr()
                trpr.append(OxmlElement("w:cantSplit"))
                if ri == 0:
                    trpr.append(OxmlElement("w:tblHeader"))
            borders = table._tbl.tblPr.find(qn("w:tblBorders"))
            if borders is None:
                borders = OxmlElement("w:tblBorders")
                table._tbl.tblPr.append(borders)
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                e = OxmlElement("w:"+edge)
                for k, v in {"val": "single", "sz": "4", "color": "D9D9D9"}.items():
                    e.set(qn("w:"+k), v)
                borders.append(e)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        else:
            doc.add_paragraph(line)
    target = ROOT / "docs/CNSRIII_四项独立研究统计分析方案_20260918_v4.docx"
    doc.core_properties.title = "CNSR III影像与长期预后四项独立研究统计分析方案"
    doc.core_properties.subject = "四项独立研究 统计方法及工作站实施"
    doc.core_properties.author = "CNSR III研究项目"
    doc.save(target)
    with (ROOT / "docs/four_studies_fields.csv").open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(["source", "canonical", "type", "permitted_codes", "unit"])
        writer.writerows((s, a, t, ";".join(map(str, c or [])), UNITS.get(a, t)) for s, (a, t, c) in SOURCES.items())
    ledger = [r for study in STUDIES for r in roles(study)]
    with (ROOT / "docs/four_studies_covariate_roles.csv").open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=list(ledger[0]))
        writer.writeheader()
        writer.writerows(ledger)
    print(target)


if __name__ == "__main__":
    main()
