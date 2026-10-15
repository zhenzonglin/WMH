"""Aggregate-only TableOne summaries and editable publication table exports.

The exporter formats previously estimated results; it never estimates a model.
Baseline summaries use the Python ``tableone`` package on the original,
unimputed, five-year-known recovery cohort. Patient identifiers are neither
selected nor copied to any output.
"""
from __future__ import annotations

import csv
import html
import re
from pathlib import Path

import numpy as np
import pandas as pd

CONTINUOUS = ("age", "wmh_ml", "gm119_ml", "icv_ml", "lesion_ml", "nihss")
CATEGORICAL = (
    "sex", "smoking", "drinking", "hypertension", "diabetes", "prior_stroke",
    "education", "pre_mrs", "toast", "mrs3",
)
PRIMARY = (
    "age", "sex", "hypertension", "diabetes", "prior_stroke", "pre_mrs",
    "nihss", "mrs3", "wmh_ml", "gm119_ml", "icv_ml", "lesion_ml",
)
SUPPLEMENT = (
    "age", "sex", "smoking", "drinking", "hypertension", "diabetes", "prior_stroke",
    "education", "pre_mrs", "toast", "nihss", "mrs3", "wmh_ml", "gm119_ml",
    "icv_ml", "lesion_ml",
)
LABELS = {
    "age": "Age, years", "sex": "Sex", "smoking": "Smoking", "drinking": "Alcohol use",
    "hypertension": "Hypertension", "diabetes": "Diabetes mellitus", "prior_stroke": "Previous stroke",
    "education": "Education", "pre_mrs": "Pre-stroke mRS score", "toast": "TOAST subtype",
    "nihss": "Admission NIHSS score", "mrs3": "Three-month mRS score",
    "wmh_ml": "WMH volume, mL", "gm119_ml": "Total grey matter volume, mL",
    "icv_ml": "Intracranial volume, mL", "lesion_ml": "Acute lesion volume, mL",
}
CODE_LABELS = {
    "sex": {1: "Male", 2: "Female"},
    "smoking": {1: "Never", 2: "Occasional", 3: "Current", 4: "Former"},
    "drinking": {1: "Never", 2: "Occasional", 3: "Current", 4: "Former"},
    "hypertension": {1: "No", 2: "Yes"},
    "diabetes": {1: "No", 2: "Yes"},
    "prior_stroke": {1: "No", 2: "Yes"},
    "education": {i: f"Education category {i}" for i in range(1, 6)},
    "pre_mrs": {i: str(i) for i in range(6)},
    "mrs3": {i: str(i) for i in range(3)},
    "toast": {1: "Large artery atherosclerosis", 2: "Cardioembolism", 3: "Small vessel occlusion",
              4: "Other determined cause", 5: "Undetermined cause"},
}
STATE_LABELS = {0: "Independent", 1: "Dependent", 2: "Dead"}
DEMO_FOOTNOTE = "SYNTHETIC STYLE PREVIEW ONLY. These simulated values are not research results."


def _column_fractions(columns: list[str]) -> list[float]:
    """Reserve width for interval estimates rather than narrow n/P columns."""
    if columns == ["Analysis", "n", "Outcome contrast", "Measure", "Effect (95% CI)", "P value"]:
        return [0.24, 0.06, 0.26, 0.09, 0.26, 0.09]
    if len(columns) == 4 and columns[0] in {"Outcome contrast", "Transition"}:
        return [0.40, 0.18, 0.30, 0.12]
    if len(columns) == 4 and columns[0] == "Scenario":
        return [0.25, 0.10, 0.20, 0.45]
    if len(columns) == 1:
        return [1.0]
    first = 0.32 if len(columns) <= 5 else 0.26
    return [first, *[(1 - first) / (len(columns) - 1)] * (len(columns) - 1)]


def _tex_escape(value: str) -> str:
    replacements = {
        "\\": r"\textbackslash{}", "&": r"\&", "%": r"\%", "$": r"\$", "#": r"\#",
        "_": r"\_", "{": r"\{", "}": r"\}", "~": r"\textasciitilde{}",
        "^": r"\textasciicircum{}", "→": r"\(\rightarrow\)", "–": "--", "—": "---", "−": "-",
        "≤": r"\(\leq\)", "≥": r"\(\geq\)", "×": r"\(\times\)", "±": r"\(\pm\)",
        "\n": " ",
    }
    return "".join(replacements.get(char, char) for char in value)


def _write_docx(path: Path, title: str, columns: list[str], rows: list[list[str]], notes: list[str]) -> None:
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Mm, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    landscape = len(columns) >= 7
    section.page_width, section.page_height = (Mm(297), Mm(210)) if landscape else (Mm(210), Mm(297))
    section.top_margin = section.bottom_margin = Mm(18)
    section.left_margin = section.right_margin = Mm(17)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.05
    title_style = document.styles["Title"]
    title_style.font.name = "Arial"
    title_style.font.size = Pt(12)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style.paragraph_format.keep_with_next = True
    title_style.paragraph_format.space_after = Pt(8)
    document.add_paragraph(title, style="Title")

    table = document.add_table(rows=1, cols=len(columns))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    available = float(section.page_width - section.left_margin - section.right_margin) / 914400
    widths = [available * fraction for fraction in _column_fractions(columns)]
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)

    # Suppress both direct and inherited rules before adding exactly three rules.
    properties = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        rule = OxmlElement(f"w:{edge}")
        rule.set(qn("w:val"), "single" if edge in {"top", "bottom"} else "nil")
        if edge in {"top", "bottom"}:
            rule.set(qn("w:sz"), "10")
            rule.set(qn("w:color"), "000000")
        borders.append(rule)
    properties.append(borders)
    cell_margins = OxmlElement("w:tblCellMar")
    for edge, value in (("top", "60"), ("bottom", "60"), ("left", "70"), ("right", "70")):
        margin = OxmlElement(f"w:{edge}")
        margin.set(qn("w:w"), value)
        margin.set(qn("w:type"), "dxa")
        cell_margins.append(margin)
    properties.append(cell_margins)

    for row_index, contents in enumerate([columns, *rows]):
        row = table.rows[0] if row_index == 0 else table.add_row()
        row_properties = row._tr.get_or_add_trPr()
        row_properties.append(OxmlElement("w:cantSplit"))
        if row_index == 0:
            row_properties.append(OxmlElement("w:tblHeader"))
        for column_index, (cell, value) in enumerate(zip(row.cells, contents)):
            cell.width = Inches(widths[column_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.space_before = Pt(2)
            paragraph.paragraph_format.keep_with_next = row_index == 0
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if column_index == 0 else WD_ALIGN_PARAGRAPH.RIGHT
            run = paragraph.add_run(value)
            run.bold = row_index == 0
            if row_index == 0:
                cell_borders = OxmlElement("w:tcBorders")
                bottom = OxmlElement("w:bottom")
                bottom.set(qn("w:val"), "single")
                bottom.set(qn("w:sz"), "6")
                bottom.set(qn("w:color"), "000000")
                cell_borders.append(bottom)
                cell._tc.get_or_add_tcPr().append(cell_borders)
        # Categorical headings should stay beside their first category.
        if row_index > 0 and all(not value for value in contents[1:]):
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = True
                    for run in paragraph.runs:
                        run.bold = True
    document.add_paragraph()
    for note in notes:
        paragraph = document.add_paragraph(note)
        paragraph.paragraph_format.space_after = Pt(3)
        for run in paragraph.runs:
            run.font.size = Pt(8)
    document.core_properties.title = title
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.comments = ""
    document.save(path)


def export_table(
    outdir: Path, stem: str, title: str, columns: list[str], rows: list[list[str]], footnotes: list[str],
    *, demo: bool = False,
) -> list[Path]:
    """Save one aggregate table as editable Word, CSV, booktabs TeX, and HTML.

    This function never accepts or processes a patient-level dataframe. The
    caller must pass aggregate rows only; cells are exported literally without
    recalculating statistics. Word tables repeat the header on continuation
    pages and prohibit a data row from splitting across pages.
    """
    if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_-]*", stem):
        raise ValueError("Table stem must contain only letters, numbers, underscores, and hyphens")
    if not columns or any(len(row) != len(columns) for row in rows):
        raise ValueError("Every aggregate row must have the same number of cells as the header")
    columns = [str(value) for value in columns]
    rows = [[str(value) for value in row] for row in rows]
    footnotes = [str(value) for value in footnotes]
    if demo:
        title = "Synthetic style preview — " + title
        footnotes = [DEMO_FOOTNOTE, *footnotes]
    outdir = Path(outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    paths = [outdir / f"{stem}{suffix}" for suffix in (".docx", ".csv", ".tex", ".html")]
    _write_docx(paths[0], title, columns, rows, footnotes)
    with paths[1].open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(columns)
        writer.writerows(rows)
        if demo:
            # Keep rectangular CSV data; metadata rows are explicitly marked.
            for marker, value in [("# Title", title), ("# Footnote", DEMO_FOOTNOTE)]:
                writer.writerow([f"{marker}: {value}", *[""] * (len(columns) - 1)])

    header = " & ".join(_tex_escape(value) for value in columns) + r" \\"
    widths = [0.90 * fraction for fraction in _column_fractions(columns)]
    specs = "".join(
        (r">{\raggedright\arraybackslash}" if i == 0 else r">{\raggedleft\arraybackslash}")
        + f"p{{{width:.3f}\\linewidth}}" for i, width in enumerate(widths)
    )
    geometry = "a4paper,margin=18mm" + (",landscape" if len(columns) >= 7 else "")
    tex = [r"\documentclass[10pt]{article}", r"\usepackage[" + geometry + r"]{geometry}",
           r"\usepackage[T1]{fontenc}", r"\usepackage[utf8]{inputenc}",
           r"\usepackage{booktabs,longtable,array}", r"\setlength{\tabcolsep}{3pt}", r"\begin{document}",
           r"\begin{longtable}{" + specs + "}", r"\caption{" + _tex_escape(title) + r"}\\",
           r"\toprule", header, r"\midrule", r"\endfirsthead", r"\toprule", header,
           r"\midrule", r"\endhead", r"\bottomrule", r"\endfoot"]
    tex.extend(" & ".join(_tex_escape(value) for value in row) + r" \\" for row in rows)
    tex.append(r"\end{longtable}")
    tex.extend(r"\noindent {\footnotesize " + _tex_escape(note) + r"}\par" for note in footnotes)
    tex.append(r"\end{document}")
    paths[2].write_text("\n".join(tex) + "\n", encoding="utf-8")

    html_rows = ["<tr>" + "".join(f"<th scope='col'>{html.escape(c)}</th>" for c in columns) + "</tr>"]
    body_rows = ["<tr>" + "".join(f"<td>{html.escape(c)}</td>" for c in row) + "</tr>" for row in rows]
    html_text = (
        "<!doctype html><html lang='en'><meta charset='utf-8'><title>" + html.escape(title) + "</title>"
        "<style>body{font:14px Arial,sans-serif;margin:32px;color:#000}h1{font-size:18px}"
        "table{border-collapse:collapse;table-layout:fixed;width:100%;border-top:1.5px solid;"
        "border-bottom:1.5px solid}"
        "thead{border-bottom:1px solid;display:table-header-group}th,td{padding:6px;text-align:right}"
        "th:first-child,td:first-child{text-align:left;white-space:pre-wrap}tr{break-inside:avoid}"
        "p{font-size:12px}th{white-space:pre-line}"
        "@page{size:A4;margin:18mm}@media print{body{margin:0;font-size:9pt}h1{font-size:12pt}"
        "th,td{padding:3pt}p{font-size:8pt}}</style><h1>" + html.escape(title) + "</h1>"
        "<table><colgroup>" + "".join(f"<col style='width:{100*f:.2f}%'>" for f in _column_fractions(columns))
        + "</colgroup><thead>" + "".join(html_rows) + "</thead><tbody>" + "".join(body_rows) + "</tbody></table>"
        + "".join(f"<p>{html.escape(note)}</p>" for note in footnotes) + "</html>"
    )
    paths[3].write_text(html_text, encoding="utf-8")
    return paths


def _numeric_source(cohort: pd.DataFrame, variables: list[str]) -> pd.DataFrame:
    missing = sorted({*variables, "state60"} - set(cohort.columns))
    if missing:
        raise ValueError("Baseline source is missing required columns: " + ", ".join(missing))
    # Whitelist rather than carrying the source frame through to TableOne.
    source = cohort.loc[:, [*variables, "state60"]].copy()
    for variable in source:
        values = pd.to_numeric(source[variable], errors="coerce")
        if (source[variable].notna() & values.isna()).any() or np.isinf(values).any():
            raise ValueError(f"Non-numeric or infinite values in baseline column {variable}")
        source[variable] = values
    if not set(source.state60.dropna().unique()).issubset(STATE_LABELS):
        raise ValueError("state60 must use 0 independent, 1 dependent, 2 dead, or missing")
    source = source.loc[source.state60.notna()].copy()
    if source.empty:
        raise ValueError("No observed five-year states are available for baseline tables")
    for variable in CONTINUOUS:
        if variable in source and source[variable].dropna().lt(0).any():
            raise ValueError(f"Negative values in baseline column {variable}")
    return source


def _baseline_content(
    cohort: pd.DataFrame, variables: list[str], education_labels: dict[str, str] | None = None
) -> tuple[list[str], list[list[str]], list[str]]:
    """Build aggregate rows with TableOne, retaining missing denominators."""
    try:
        from tableone import TableOne
    except ImportError as error:
        raise ImportError("Install publication dependencies with: python -m pip install "
                          "-c requirements-conda.txt -r requirements-publication.txt") from error

    source = _numeric_source(cohort, variables)
    labels = {variable: dict(values) for variable, values in CODE_LABELS.items()}
    education_unverified = "education" in variables and education_labels is None
    if education_labels is not None:
        if not isinstance(education_labels, dict):
            raise ValueError("Education labels must be a JSON object mapping codes to verified labels")
        try:
            labels["education"] = {int(code): str(label).strip() for code, label in education_labels.items()}
        except (ValueError, TypeError) as error:
            raise ValueError("Education label keys must be integer category codes") from error
        if any(not value for value in labels["education"].values()):
            raise ValueError("Education category labels cannot be blank")
        if len(set(labels["education"].values())) != len(labels["education"]):
            raise ValueError("Education category labels must be unique")

    categorical = [variable for variable in variables if variable in CATEGORICAL]
    display = source.copy()
    for variable in categorical:
        observed = set(source[variable].dropna().unique())
        if not observed.issubset(labels[variable]):
            raise ValueError(f"Unmapped category codes in baseline column {variable}")
        display[variable] = source[variable].map(labels[variable])
    display["state60"] = source.state60.map(STATE_LABELS)
    tab = TableOne(
        display, columns=variables, categorical=categorical,
        nonnormal=[variable for variable in variables if variable in CONTINUOUS],
        groupby="state60", overall=True, pval=False, missing=True, row_percent=False, include_null=False,
        decimals=1, order={variable: [label for code, label in labels[variable].items()
                                    if source[variable].eq(code).any()] for variable in categorical},
    ).tableone
    group_names = ["Overall", *STATE_LABELS.values()]
    actual_columns = {str(column[-1] if isinstance(column, tuple) else column): column for column in tab.columns}
    available_groups = set(source.state60.unique())
    counts = [len(source), *(int(source.state60.eq(state).sum()) for state in STATE_LABELS)]
    columns = ["Characteristic", *(
        f"{name}\n(n={count:,})" for name, count in zip(["All known", *STATE_LABELS.values()], counts)
    )]
    cells: dict[str, list[tuple[str, list[str]]]] = {variable: [] for variable in variables}
    for index, row in tab.iterrows():
        field = str(index[0] if isinstance(index, tuple) else index).split(",", 1)[0]
        if field not in cells:
            continue
        category = str(index[1]) if isinstance(index, tuple) and len(index) > 1 else ""
        values = [str(row[actual_columns[group]]) if group in actual_columns else "Not observed"
                  for group in group_names]
        cells[field].append((category, values))
    rows: list[list[str]] = []
    for variable in variables:
        entries = cells[variable]
        if not entries:
            raise RuntimeError(f"TableOne returned no summary for {variable}")
        if variable in categorical:
            rows.append([LABELS[variable], "", "", "", ""])
            entry_map = dict(entries)
            for category in labels[variable].values():
                if category in entry_map:
                    rows.append([f"  {category}", *entry_map[category]])
        else:
            if len(entries) != 1:
                raise RuntimeError(f"Unexpected TableOne continuous summary structure for {variable}")
            rows.append([LABELS[variable], *entries[0][1]])
        missing_counts = [int(source[variable].isna().sum()), *(
            int(source.loc[source.state60.eq(state), variable].isna().sum()) for state in STATE_LABELS
        )]
        if any(missing_counts):
            rows.append(["  Missing, n", *(str(count) for count in missing_counts)])
    footnotes = [
        ("Data are median [25th, 75th percentile] or number (percentage). Summaries were calculated with "
         "Python TableOne from observed values before covariate imputation."),
        ("Only participants with an observed five-year functional state are included. Independent: mRS 0-2; "
        "dependent: mRS 3-5; dead: confirmed death. Percentages use nonmissing observations within each "
        "column. Missing values are shown as counts; variables without a Missing row have no missing values."),
        ("No baseline significance tests were performed. WMH, white matter hyperintensity; mRS, modified "
        "Rankin Scale; NIHSS, National Institutes of Health Stroke Scale; TOAST, Trial of Org 10172 in Acute "
        "Stroke Treatment."),
    ]
    if available_groups != set(STATE_LABELS):
        footnotes.append("Not observed denotes a functional-state group containing no participants.")
    if education_unverified:
        footnotes.append("Education category names are pending verification against the source data dictionary; "
                         "category codes have been retained without assigning unverified educational levels.")
    return columns, rows, footnotes


def build_baseline_tables(
    cohort: pd.DataFrame, outdir: Path, education_labels: dict[str, str] | None = None,
    *, demo: bool = False,
) -> list[Path]:
    """Export main Table 1 and full supplementary Table 1 using genuine TableOne.

    ``cohort`` must contain all requested original covariates and ``state60``
    coded 0/1/2/NaN. Five-year-unknown participants are excluded only from these
    five-year-stratified descriptive tables. No analysis cohort is modified.
    """
    prepared: list[tuple[str, str, list[str], list[list[str]], list[str]]] = []
    for stem, title, variables in (
        ("table1", "Table 1 Baseline characteristics by five-year functional state", PRIMARY),
        ("supplement_table1", "Supplementary Table 1 Full baseline characteristics by five-year functional state",
         SUPPLEMENT),
    ):
        columns, rows, notes = _baseline_content(cohort, list(variables), education_labels)
        prepared.append((stem, title, columns, rows, notes))
    # Validate both tables before writing any output, including source codebooks.
    results: list[Path] = []
    for stem, title, columns, rows, notes in prepared:
        results.extend(export_table(Path(outdir), stem, title, columns, rows, notes, demo=demo))
    return results
