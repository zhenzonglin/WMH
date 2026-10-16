"""Revise one descriptive baseline row from exact exported integer frequencies.

This utility never constructs patient records, runs TableOne, imputes values,
or fits a model. It evaluates the same linear-interpolated order statistics
as NumPy/TableOne directly from cumulative category counts. Future raw-data
exports use the genuine TableOne implementation in ``tables.py``.
"""
from __future__ import annotations

import csv
import hashlib
import math
import re
import shutil
import tempfile
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile

from .tables import export_table

STEMS = ("table1", "supplement_table1")
FORMATS = (".csv", ".docx", ".html", ".tex")
ROW_LABEL = "Pre-stroke mRS score"
REVISION_NOTE = (
    "Pre-stroke mRS quantiles were calculated from the complete observed category frequencies."
)
_COUNT_CELL = re.compile(r"^\s*([0-9][0-9,]*)\s*\(\s*([0-9]+(?:\.[0-9]+)?)\s*%?\s*\)\s*$")
_CONTINUOUS_CELL = re.compile(
    r"^\s*([0-9]+(?:\.[0-9]+)?)\s*\[\s*([0-9]+(?:\.[0-9]+)?),\s*([0-9]+(?:\.[0-9]+)?)\s*\]\s*$"
)


def linear_quantiles_from_counts(counts: dict[int, int], probabilities=(.25, .5, .75)) -> list[float]:
    """Exact NumPy ``method='linear'`` quantiles without expanding observations."""
    if not counts or any(isinstance(k, bool) or not isinstance(k, int) or k not in range(6) for k in counts):
        raise ValueError("Pre-stroke mRS frequency keys must be integer scores 0 through 5")
    if any(isinstance(n, bool) or not isinstance(n, int) or n < 0 for n in counts.values()):
        raise ValueError("Frequency counts must be nonnegative integers")
    n = sum(counts.values())
    if not n:
        raise ValueError("A quantile requires at least one observed score")
    if any(not math.isfinite(q) or not 0 <= q <= 1 for q in probabilities):
        raise ValueError("Quantile probabilities must lie within [0,1]")

    def at(position: int) -> int:
        cumulative = 0
        for score, count in sorted(counts.items()):
            cumulative += count
            if position < cumulative:
                return score
        raise ValueError("Order statistic falls outside the frequency distribution")

    values = []
    for probability in probabilities:
        position = (n - 1) * probability
        lower, upper = math.floor(position), math.ceil(position)
        fraction = position - lower
        values.append(float(at(lower) + fraction * (at(upper) - at(lower))))
    return values


def _denominators(columns: list[str]) -> list[int]:
    names = ("All known", "Independent", "Dependent", "Dead")
    if len(columns) != 5 or columns[0] != "Characteristic":
        raise ValueError("Expected the five-column published recovery baseline table")
    counts = []
    for column, name in zip(columns[1:], names, strict=True):
        label = r"(?:All known|All|Overall)" if name == "All known" else re.escape(name)
        match = re.fullmatch(rf"{label}\s*\(n=([0-9][0-9,]*)\)", column.strip())
        if match is None:
            raise ValueError("Baseline columns must retain the original cohort labels and denominators")
        counts.append(int(match[1].replace(",", "")))
    if counts[0] != sum(counts[1:]):
        raise ValueError("Overall baseline N must equal the sum of the three outcome-group counts")
    return counts


def _revise_rows(columns: list[str], rows: list[list[str]]) -> tuple[list[list[str]], dict]:
    denominators = _denominators(columns)
    matches = [index for index, row in enumerate(rows) if row[0].strip() == ROW_LABEL]
    if len(matches) != 1:
        raise ValueError("Expected exactly one pre-stroke mRS row")
    start = matches[0]
    if any(rows[start][1:]):
        cells = [_CONTINUOUS_CELL.fullmatch(cell) for cell in rows[start][1:]]
        if not all(cells):
            raise ValueError("Existing pre-stroke row is neither a categorical block nor median [Q1,Q3]")
        for match in cells:
            median, q1, q3 = map(float, match.groups())
            if not 0 <= q1 <= median <= q3 <= 5:
                raise ValueError("Existing pre-stroke median/quartiles are outside the valid score range")
        return rows, {"status": "ALREADY_CONTINUOUS", "denominators": denominators,
                      "replacement_row": rows[start], "changed": False}
    end = start + 1
    counts: dict[int, list[int]] = {}
    percentages: dict[int, list[float]] = {}
    missing = [0] * 4
    missing_row = None
    while end < len(rows) and rows[end][0].startswith(" "):
        row = rows[end]
        label = row[0].strip()
        if label == "Missing, n":
            if missing_row is not None or any(not re.fullmatch(r"[0-9]+", value) for value in row[1:]):
                raise ValueError("Malformed or duplicated pre-stroke mRS missing-count row")
            missing = [int(value) for value in row[1:]]
            missing_row = row
        else:
            if not re.fullmatch(r"[0-5]", label) or int(label) in counts or missing_row is not None:
                raise ValueError("Unexpected or duplicated category within the pre-stroke mRS block")
            matches = [_COUNT_CELL.fullmatch(cell) for cell in row[1:]]
            if not all(matches):
                raise ValueError("Pre-stroke mRS requires exact integer counts, not percentages alone")
            counts[int(label)] = [int(match[1].replace(",", "")) for match in matches]
            percentages[int(label)] = [float(match[2]) for match in matches]
        end += 1
    if not counts:
        raise ValueError("No exact category counts are available for pre-stroke mRS")
    for row in [*counts.values(), missing]:
        if row[0] != sum(row[1:]):
            raise ValueError("Pre-stroke overall counts differ from the sum of the outcome groups")
    observed = [sum(values[column] for values in counts.values()) for column in range(4)]
    if any(n + absent != total for n, absent, total in zip(observed, missing, denominators, strict=True)):
        raise ValueError("Pre-stroke category and missing counts do not reconcile with baseline N")
    if not all(observed):
        raise ValueError("No observed pre-stroke scores in one or more baseline columns")
    for score, values in counts.items():
        for column, count in enumerate(values):
            if abs(100 * count / observed[column] - percentages[score][column]) > .051:
                raise ValueError("Pre-stroke displayed percentages disagree with exact category counts")
    quantiles = [linear_quantiles_from_counts({score: values[col] for score, values in counts.items()})
                 for col in range(4)]
    replacement = [ROW_LABEL, *(f"{median:.1f} [{q1:.1f},{q3:.1f}]" for q1, median, q3 in quantiles)]
    new_block = [replacement] + ([missing_row] if missing_row is not None else [])
    revised = [*rows[:start], *new_block, *rows[end:]]
    return revised, {
        "status": "REVISED_FROM_EXACT_COUNTS", "changed": True,
        "denominators": denominators, "observed_n": observed, "missing_n": missing,
        "category_counts": {str(score): values for score, values in sorted(counts.items())},
        "column_order": ["All known", "Independent", "Dependent", "Dead"],
        "quantiles_q25_q50_q75": quantiles, "replacement_row": replacement,
        "original_block_rows": end - start, "replacement_block_rows": len(new_block),
        "source_word_row_index": start + 1,
        "unmodified_rows": len(rows) - (end - start),
        "method": "Linear order statistics: h=(n-1)*q; interpolate order floor(h) and ceil(h).",
        "patient_records_reconstructed": False, "model_refitted": False,
    }


def _read_export(source: Path, stem: str) -> tuple[str, list[str], list[list[str]], list[str], dict]:
    from docx import Document

    evidence = {suffix: {
        "name": stem + suffix,
        "sha256": hashlib.sha256((source / (stem + suffix)).read_bytes()).hexdigest(),
    } for suffix in FORMATS}
    with (source / f"{stem}.csv").open(encoding="utf-8-sig", newline="") as handle:
        content = [row for row in csv.reader(handle) if row]
    if not content or any(len(row) != 5 for row in content):
        raise ValueError(f"{stem}: expected a rectangular baseline CSV")
    columns, rows = content[0], content[1:]
    document = Document(source / f"{stem}.docx")
    if len(document.tables) != 1:
        raise ValueError(f"{stem}: Word export must contain exactly one table")
    word_rows = [[cell.text for cell in row.cells] for row in document.tables[0].rows]

    def normalized(table):
        return [[cell.replace("\r", "").strip() for cell in row] for row in table]

    compared_word, compared_csv = normalized(word_rows), normalized(content)
    alias_preserved = False
    if len(compared_word) == len(compared_csv) and all(len(row) == 5 for row in compared_word):
        word_header, csv_header = compared_word[0][1], compared_csv[0][1]
        pattern = r"(?:All known|All|Overall)\s*\(n=([0-9][0-9,]*)\)"
        word_match, csv_match = re.fullmatch(pattern, word_header), re.fullmatch(pattern, csv_header)
        if (word_match and csv_match
                and int(word_match[1].replace(",", "")) == int(csv_match[1].replace(",", ""))):
            alias_preserved = word_header != csv_header
            compared_word[0][1] = csv_header
    if compared_word != compared_csv:
        raise ValueError(f"{stem}: Word and CSV contents disagree; reconcile the source exports first")
    source_csv_columns = list(columns)
    columns = [cell.replace("\r", "") for cell in word_rows[0]]
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    if not paragraphs:
        raise ValueError(f"{stem}: Word export lacks a title")
    new_rows, audit = _revise_rows(columns, rows)
    notes = paragraphs[1:]
    if audit["changed"] and REVISION_NOTE not in notes:
        notes = [*notes, REVISION_NOTE]
    audit["source_files"] = evidence
    audit["source_csv_columns"] = source_csv_columns
    audit["preserved_word_columns"] = columns
    audit["overall_header_alias_preserved"] = alias_preserved
    return paragraphs[0], columns, new_rows, notes, audit


def _patch_word_row(source: Path, destination: Path, audit: dict) -> None:
    """Change only document.xml; preserve all other original Word package parts."""
    from lxml import etree

    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    def tag(name):
        return "{" + ns["w"] + "}" + name
    with ZipFile(source) as archive:
        tree = etree.fromstring(archive.read("word/document.xml"))
        body = tree.find(tag("body"))
        tables = body.findall(tag("tbl"))
        if len(tables) != 1:
            raise ValueError("Targeted Word revision requires one top-level baseline table")
        table = tables[0]
        rows = table.findall(tag("tr"))
        index = audit["source_word_row_index"]
        target = rows[index]
        cells = target.findall(tag("tc"))
        def text(element):
            return "".join(element.itertext())
        if len(cells) != 5 or text(cells[0]).strip() != ROW_LABEL:
            raise ValueError("Targeted Word row no longer matches pre-stroke mRS")
        template_cells = rows[1].findall(tag("tc"))  # Existing continuous Age row, same columns.

        def replace_paragraph(paragraph, value, template):
            properties = template.find(tag("pPr"))
            run_properties = template.find("w:r/w:rPr", ns)
            for child in list(paragraph):
                paragraph.remove(child)
            if properties is not None:
                paragraph.append(deepcopy(properties))
            run = etree.SubElement(paragraph, tag("r"))
            if run_properties is not None:
                run.append(deepcopy(run_properties))
            etree.SubElement(run, tag("t")).text = value

        for cell, template_cell, value in zip(cells, template_cells, audit["replacement_row"], strict=True):
            paragraph = cell.find(tag("p"))
            template = template_cell.find(tag("p"))
            if paragraph is None or template is None:
                raise ValueError("Word baseline cell lacks its original paragraph")
            replace_paragraph(paragraph, value, template)
            for extra in cell.findall(tag("p"))[1:]:
                cell.remove(extra)
        for row in rows[index + 1:index + audit["original_block_rows"]]:
            if text(row.findall(tag("tc"))[0]).strip() != "Missing, n":
                table.remove(row)
        paragraphs = [paragraph for paragraph in body.findall(tag("p")) if text(paragraph).strip()]
        if not any(text(paragraph) == REVISION_NOTE for paragraph in paragraphs):
            paragraph = etree.Element(tag("p"))
            replace_paragraph(paragraph, REVISION_NOTE, paragraphs[-1])
            section = body.find(tag("sectPr"))
            if section is None:
                body.append(paragraph)
            else:
                section.addprevious(paragraph)
        payload = etree.tostring(tree, encoding="UTF-8", xml_declaration=True, standalone=True)
        with ZipFile(destination, "w") as result:
            for info in archive.infolist():
                result.writestr(info, payload if info.filename == "word/document.xml" else archive.read(info.filename))


def revise_baseline_exports(source_table_dir: Path, dest_table_dir: Path) -> dict:
    """Revise only Table 1/Supplementary Table 1 in a separate copied package.

    The destination may already contain a package copy; only these two stems
    and their four export formats are replaced. Both tables are validated
    before writing. Already-continuous exports are copied byte-for-byte.
    The returned audit uses JSON-serializable values only.
    """
    source, destination = Path(source_table_dir).resolve(), Path(dest_table_dir).resolve()
    if source == destination or source in destination.parents or destination in source.parents:
        raise ValueError("Baseline revision destination must be separate from the source table directory")
    prepared = {stem: _read_export(source, stem) for stem in STEMS}
    audits = {stem: contents[-1] for stem, contents in prepared.items()}
    if audits[STEMS[0]]["replacement_row"] != audits[STEMS[1]]["replacement_row"]:
        raise ValueError("Main and supplementary pre-stroke mRS summaries disagree")
    if audits[STEMS[0]].get("category_counts") != audits[STEMS[1]].get("category_counts"):
        raise ValueError("Main and supplementary pre-stroke mRS frequency evidence disagrees")
    for stem, audit in audits.items():
        for suffix, evidence in audit["source_files"].items():
            if hashlib.sha256((source / (stem + suffix)).read_bytes()).hexdigest() != evidence["sha256"]:
                raise ValueError("Source baseline export changed while reading")
    destination.mkdir(parents=True, exist_ok=True)
    outputs = []
    for stem, (title, columns, rows, notes, audit) in prepared.items():
        if audit["changed"]:
            # Only non-Word formats use the generic exporter. Preserve the
            # author's existing Word styles/layout with a targeted OOXML edit.
            with tempfile.TemporaryDirectory(prefix="baseline_table_export_", dir=destination.parent) as temp:
                export_table(Path(temp), stem, title, columns, rows, notes)
                for suffix in FORMATS:
                    target = destination / (stem + suffix)
                    if suffix == ".docx":
                        _patch_word_row(source / (stem + suffix), target, audit)
                    else:
                        shutil.copy2(Path(temp) / (stem + suffix), target)
                    outputs.append(str(target))
        else:
            for suffix in FORMATS:
                target = destination / (stem + suffix)
                shutil.copy2(source / (stem + suffix), target)
                outputs.append(str(target))
    for stem, audit in audits.items():
        for suffix, evidence in audit["source_files"].items():
            if hashlib.sha256((source / (stem + suffix)).read_bytes()).hexdigest() != evidence["sha256"]:
                raise ValueError("Source baseline export changed during revision")
    return {"status": "REVISED" if any(a["changed"] for a in audits.values()) else "NO_CHANGE",
            "tables": audits, "outputs": outputs, "other_statistics_changed": False,
            "patient_records_reconstructed": False, "model_refitted": False}
