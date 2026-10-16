"""Exact aggregate quantiles and isolated replacement of one descriptive row."""
from __future__ import annotations

import csv
import hashlib
import json
from zipfile import ZipFile

import numpy as np
import pytest
from docx import Document

from wmh_hcy.recovery_publication.baseline_revision import (
    FORMATS,
    REVISION_NOTE,
    STEMS,
    _revise_rows,
    linear_quantiles_from_counts,
    revise_baseline_exports,
)
from wmh_hcy.recovery_publication.tables import export_table


def _content(missing=False):
    totals = (10, 5, 3, 2) if missing else (8, 4, 2, 2)
    columns = ["Characteristic", *(f"{name}\n(n={n})" for name, n in zip(
        ("All known", "Independent", "Dependent", "Dead"), totals, strict=True))]
    rows = [["Age, years", "60.0 [50.0,70.0]", "60.0 [50.0,70.0]", "60.0 [50.0,70.0]", "60.0 [50.0,70.0]"],
            ["Pre-stroke mRS score", "", "", "", ""],
            ["  0", "4 (50.0)", "3 (75.0)", "1 (50.0)", "0 (0.0)"],
            ["  1", "2 (25.0)", "1 (25.0)", "1 (50.0)", "0 (0.0)"],
            ["  2", "2 (25.0)", "0 (0.0)", "0 (0.0)", "2 (100.0)"]]
    if missing:
        rows.append(["  Missing, n", "2", "1", "1", "0"])
    rows += [["Three-month mRS score", "", "", "", ""],
             ["  0", "4 (50.0)", "2 (50.0)", "1 (50.0)", "1 (50.0)"],
             ["  1", "4 (50.0)", "2 (50.0)", "1 (50.0)", "1 (50.0)"]]
    return columns, rows


@pytest.mark.parametrize("counts", [
    {0: 3, 1: 1}, {0: 1, 5: 1}, {1: 1}, {0: 2, 1: 3, 2: 1, 3: 2, 4: 1, 5: 1},
    {0: 0, 1: 2, 2: 2},
])
def test_frequency_quantiles_match_numpy_linear_order_statistics(counts):
    # Expansion is confined to tiny synthetic test fixtures, never production.
    raw = np.repeat(list(counts), list(counts.values()))
    np.testing.assert_allclose(linear_quantiles_from_counts(counts), np.quantile(raw, [.25, .5, .75]))


@pytest.mark.parametrize("counts", [{0: -1}, {0: 1.5}, {0: 0}, {6: 1}, {True: 1}])
def test_invalid_frequency_inputs_are_rejected(counts):
    with pytest.raises(ValueError):
        linear_quantiles_from_counts(counts)


def test_only_prestroke_block_changes_and_missing_counts_remain():
    columns, rows = _content(missing=True)
    result, audit = _revise_rows(columns, rows)
    assert result[0] == rows[0]
    assert result[1] == ["Pre-stroke mRS score", "0.5 [0.0,1.2]", "0.0 [0.0,0.2]",
                         "0.5 [0.2,0.8]", "2.0 [2.0,2.0]"]
    assert result[2] == ["  Missing, n", "2", "1", "1", "0"]
    assert result[3:] == rows[6:]
    assert audit["missing_n"] == [2, 1, 1, 0]
    assert audit["observed_n"] == [8, 4, 2, 2]
    assert audit["patient_records_reconstructed"] is False


@pytest.mark.parametrize("mutation", ["denominator", "overall", "fraction", "percent_only", "percent"])
def test_incomplete_or_inconsistent_counts_block_revision(mutation):
    columns, rows = _content()
    if mutation == "denominator":
        columns[1] = "All known\n(n=9)"
    elif mutation == "overall":
        rows[2][1] = "3 (50.0)"
    elif mutation == "fraction":
        rows[2][1] = "4.0 (50.0)"
    elif mutation == "percent_only":
        rows[2][1] = "50.0%"
    else:
        rows[2][1] = "4 (80.0)"
    with pytest.raises(ValueError):
        _revise_rows(columns, rows)


def test_reexport_keeps_titles_notes_other_rows_and_sources(tmp_path):
    source, target = tmp_path / "original", tmp_path / "revision"
    columns, rows = _content()
    for stem in STEMS:
        export_table(source, stem, f"Original {stem} title", columns, rows, ["Original TableOne footnote."])
    before = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in source.iterdir()}
    audit = revise_baseline_exports(source, target)
    assert audit["status"] == "REVISED" and len(audit["outputs"]) == 8
    json.dumps(audit, allow_nan=False)
    assert before == {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in source.iterdir()}
    for stem in STEMS:
        doc = Document(target / f"{stem}.docx")
        text = [p.text for p in doc.paragraphs]
        assert text[0] == f"Original {stem} title"
        assert "Original TableOne footnote." in text and REVISION_NOTE in text
        with (target / f"{stem}.csv").open(encoding="utf-8-sig", newline="") as handle:
            revised = list(csv.reader(handle))
        assert revised[0] == columns
        assert revised[1] == rows[0]
        assert revised[3:] == rows[5:]
        with ZipFile(source / f"{stem}.docx") as original, ZipFile(target / f"{stem}.docx") as edited:
            assert original.namelist() == edited.namelist()
            for name in original.namelist():
                if name != "word/document.xml":
                    assert original.read(name) == edited.read(name)
    repeated = tmp_path / "repeated"
    no_change = revise_baseline_exports(target, repeated)
    assert no_change["status"] == "NO_CHANGE"
    for stem in STEMS:
        for suffix in FORMATS:
            assert (target / (stem + suffix)).read_bytes() == (repeated / (stem + suffix)).read_bytes()


def test_word_csv_disagreement_blocks_before_output(tmp_path):
    columns, rows = _content()
    source, target = tmp_path / "source", tmp_path / "target"
    for stem in STEMS:
        export_table(source, stem, "Title", columns, rows, [])
    doc = Document(source / "table1.docx")
    doc.tables[0].cell(1, 1).text = "Different saved value"
    doc.save(source / "table1.docx")
    with pytest.raises(ValueError, match="Word and CSV contents disagree"):
        revise_baseline_exports(source, target)
    assert not target.exists()


def test_only_equivalent_overall_header_alias_is_preserved(tmp_path):
    columns, rows = _content()
    source, target = tmp_path / "source", tmp_path / "target"
    for stem in STEMS:
        export_table(source, stem, "Title", columns, rows, [])
    doc = Document(source / "table1.docx")
    doc.tables[0].cell(0, 1).text = "All \n(n=8)"
    doc.save(source / "table1.docx")
    audit = revise_baseline_exports(source, target)
    assert audit["tables"]["table1"]["overall_header_alias_preserved"]
    with (target / "table1.csv").open(encoding="utf-8-sig", newline="") as handle:
        header = next(csv.reader(handle))
    assert header[1] == "All \n(n=8)"
    assert Document(target / "table1.docx").tables[0].cell(0, 1).text == header[1]
    doc.tables[0].cell(0, 1).text = "All \n(n=9)"
    doc.save(source / "table1.docx")
    with pytest.raises(ValueError, match="Word and CSV contents disagree"):
        revise_baseline_exports(source, tmp_path / "invalid")


def test_source_destination_overlap_is_rejected(tmp_path):
    with pytest.raises(ValueError, match="separate"):
        revise_baseline_exports(tmp_path, tmp_path)
